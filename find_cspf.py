import openpyxl

def find_cspf_logic(file_path, sheet_name):
    wb = openpyxl.load_workbook(file_path, data_only=False)
    sheet = wb[sheet_name]
    
    # Search for CSPF keyword
    target_cells = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and "CSPF" in cell.value:
                target_cells.append(cell)
    
    for cell in target_cells:
        print(f"Found keyword at {cell.coordinate}: {cell.value}")
        # Print surrounding cells
        r, c = cell.row, cell.column
        for i in range(max(1, r-2), r+5):
            row_data = []
            for j in range(max(1, c-2), c+10):
                adj_cell = sheet.cell(row=i, column=j)
                row_data.append(f"[{adj_cell.coordinate}]: {adj_cell.value}")
            print(" | ".join(row_data))
        print("-" * 20)

if __name__ == "__main__":
    find_cspf_logic('20181107 ISO16358-1_AMD1 Calculation_tool_FINAL (1).xlsm', 'Variable Capacity unit')
