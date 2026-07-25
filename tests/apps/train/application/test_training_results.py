"""Phase 5C structured result, report, and lifecycle invariants."""

from __future__ import annotations

import csv
import json
from dataclasses import replace
import pytest
from openpyxl import load_workbook
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

from apps.train.adapters.training_results import (
    TrainingResultArtifactWriter,
)
from apps.train.application.training_results import (
    TrainingAnalysisResult,
    TrainingResultService,
    load_training_analysis_payload,
)
from core.ml.training_results import CoreTrainingEvidence, summarize_fold_metrics


def _target(identity: str, offset: float = 0.0) -> dict:
    return {
        "target_identity": identity,
        "target_ml_name": f"ml_{identity}",
        "status": "complete",
        "metrics": {
            "evaluation_scope": "shuffled_5_fold_cross_validation",
            "sample_count": 10,
            "fold_count": 5,
            "seed": 42,
            "r2": 0.8 + offset,
            "mae": 1.5 - offset,
            "rmse": 2.0 - offset,
            "r2_std": 0.01,
            "mae_std": 0.02,
            "rmse_std": 0.03,
        },
        "rfecv": {
            "status": "used",
            "score_context": "neg_root_mean_squared_error; 5-fold CV",
            "feature_count_before": 2,
            "feature_count_after": 1,
            "features": [
                {
                    "feature": "f1",
                    "selected": True,
                    "rank": 1,
                    "selection_order": None,
                },
                {
                    "feature": "f2",
                    "selected": False,
                    "rank": 2,
                    "selection_order": None,
                },
            ],
        },
        "feature_importance": [{
            "feature": "f1",
            "method": "xgboost_gain",
            "raw_value": 0.75,
            "normalized_value": 1.0,
            "rank": 1,
        }],
        "optuna": {
            "status": "used",
            "score_context": "mean 5-fold CV RMSE",
            "selected_parameters": {"max_depth": 4},
            "trials": [{
                "trial_number": 0,
                "parameters": {"max_depth": 4},
                "score": 2.0,
                "state": "COMPLETE",
                "duration_seconds": 0.5,
                "best_trial": True,
            }],
        },
    }


def _evidence(*targets: dict, status: str = "complete") -> CoreTrainingEvidence:
    return CoreTrainingEvidence(
        started_at="2026-07-25T00:00:00+00:00",
        finished_at="2026-07-25T00:00:10+00:00",
        duration_seconds=10.0,
        training_data_sha256="data-sha",
        training_data_rows=10,
        evaluation_context={
            "scope": "shuffled_5_fold_cross_validation",
            "fold_count": 5,
            "seed": 42,
            "splitter": "KFold",
        },
        targets=targets,
        preprocessing={
            "status": "applied",
            "steps": [{"name": "drop_missing_base_features"}],
            "feature_data_quality": [{
                "feature": "f1",
                "missing_count": 1,
                "missing_rate": 0.1,
                "unique_count": 9,
                "low_variance": False,
                "outlier_method": "iqr_1.5",
                "outlier_count": 1,
                "target_usage": ("cooling", "heating"),
                "selection_state": "before_target_policy",
            }],
        },
        status=status,
    )


def test_metrics_match_independent_known_prediction_recalculation():
    actual = ([1.0, 2.0, 3.0], [1.0, 2.5, 2.0])
    second = ([2.0, 4.0, 6.0], [2.5, 3.5, 5.5])
    folds = []
    for truth, prediction in (actual, second):
        folds.append({
            "r2": r2_score(truth, prediction),
            "mae": mean_absolute_error(truth, prediction),
            "rmse": mean_squared_error(truth, prediction) ** 0.5,
        })

    metrics = summarize_fold_metrics(folds, sample_count=6)

    assert metrics["r2"] == pytest.approx(sum(x["r2"] for x in folds) / 2)
    assert metrics["mae"] == pytest.approx(sum(x["mae"] for x in folds) / 2)
    assert metrics["rmse"] == pytest.approx(sum(x["rmse"] for x in folds) / 2)
    assert metrics["sample_count"] == 6


def test_fair_baseline_deltas_and_invalid_context_suppression():
    service = TrainingResultService()
    baseline = service.build(
        run_id="baseline-run",
        candidate_id="baseline",
        evidence=_evidence(_target("cooling"), _target("heating")),
        required_target_identities=("cooling", "heating"),
        publication_outcome="published",
    )
    current = service.build(
        run_id="current-run",
        candidate_id="current",
        evidence=_evidence(_target("cooling", 0.1), _target("heating", 0.05)),
        required_target_identities=("cooling", "heating"),
        baseline=baseline,
    )

    assert current.baseline["comparable"] is True
    assert current.targets[0]["target_identity"] == "cooling"
    assert current.targets[0]["baseline_comparison"]["delta"]["r2"] == pytest.approx(0.1)

    unfair = service.build(
        run_id="unfair",
        candidate_id="unfair",
        evidence=replace(
            _evidence(_target("cooling"), _target("heating")),
            training_data_sha256="different",
        ),
        required_target_identities=("cooling", "heating"),
        baseline=baseline,
    )
    assert unfair.baseline["comparable"] is False
    assert "data context" in unfair.baseline["unavailable_reason"]
    assert "delta" not in unfair.targets[0]["baseline_comparison"]


def test_partial_is_non_promotable_and_preserves_successful_target():
    failed = {
        "target_identity": "heating",
        "target_ml_name": "ml_heating",
        "status": "failed",
        "blocking_reason": "fit failed",
    }
    result = TrainingResultService().build(
        run_id="partial",
        candidate_id="candidate-partial",
        evidence=_evidence(_target("cooling"), failed, status="partial"),
        required_target_identities=("cooling", "heating"),
    )

    assert result.run["status"] == "partial"
    assert result.targets[0]["metrics"]["r2"] == 0.8
    assert result.promotion_eligibility["eligible"] is False
    assert result.promotion_eligibility["snapshot_only"] is True


def test_json_csv_and_xlsx_derive_from_same_contract(tmp_path):
    result = TrainingResultService().build(
        run_id="run",
        candidate_id="candidate",
        evidence=_evidence(_target("cooling"), _target("heating")),
        required_target_identities=("cooling", "heating"),
    )
    TrainingResultArtifactWriter().write(tmp_path, result)

    payload = json.loads((tmp_path / "training_result.json").read_text())
    with (tmp_path / "analysis" / "target_metrics.csv").open(
        newline="", encoding="utf-8"
    ) as source:
        metrics = list(csv.DictReader(source))
    workbook = load_workbook(tmp_path / "training_report.xlsx", data_only=True)

    assert payload["targets"][0]["metrics"]["r2"] == float(metrics[0]["r2"])
    assert workbook["Target Metrics"]["F2"].value == payload["targets"][0]["metrics"]["r2"]
    assert workbook.sheetnames == [
        "Summary", "Target Metrics", "Selected Features", "RFECV Ranking",
        "Feature Importance", "Optuna Best Parameters", "Optuna Trials",
        "Preprocessing", "Run Information",
    ]


def test_legacy_unavailable_and_future_version_rejected():
    assert load_training_analysis_payload(None)["status"] == "unavailable"
    with pytest.raises(ValueError, match="unsupported training result"):
        load_training_analysis_payload({"schema_version": "training_result.v999"})
