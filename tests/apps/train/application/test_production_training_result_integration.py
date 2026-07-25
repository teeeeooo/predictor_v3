"""Bounded production-owner regression from training data to Candidate."""

from __future__ import annotations

import csv
import json
import os
from pathlib import Path

from openpyxl import load_workbook

from apps.common.model_lifecycle import ModelLifecycleRepository
from apps.train.application.training_lifecycle import TrainingLifecycleService
from apps.train.composition.training_results import build_candidate_publisher
from apps.train.jobs.train_job import run_production_training
from apps.train.state.training_run_state import TrainingRequest, TrainingResult
from core.data_definition.contract import bootstrap_manifest
from core.data_definition.target_registry.runtime import (
    apply_target_policy,
    model_registry_snapshot,
)
from core.ml.preprocessing import load_and_preprocess, prepare_pipeline
from core.ml.training_results import TrainingOptimizationConfig
from tests.apps.common.model_lifecycle.conftest import publish_candidate
from tools.dev.mock_smoke.generators import write_mock_training_data


class _BoundedProductionExecution:
    def __init__(self, snapshot) -> None:  # noqa: ANN001
        self._snapshot = snapshot

    @property
    def is_running(self) -> bool:
        return False

    def start(self, request, callbacks) -> None:  # noqa: ANN001
        final_path = Path(request.model_output_path)
        temporary = final_path.with_name(".production-training.tmp")
        output, _evidence_path = run_production_training(
            request,
            temporary,
            log_callback=lambda _message: None,
            registry_snapshot=self._snapshot,
            optimization_config=TrainingOptimizationConfig(
                cv_folds=2,
                optuna_trials=1,
                n_estimators_min=2,
                n_estimators_max=2,
                n_jobs=1,
                optuna_sampler_seed=42,
            ),
        )
        os.replace(temporary, final_path)
        callbacks.finished(TrainingResult(
            request.run_id,
            output.evidence.status,
            summary=output.summary,
            model_path=str(final_path),
        ))

    def cancel(self) -> bool:
        return False

    def dispose(self) -> None:
        return None


def test_production_multi_target_training_publishes_complete_analysis(tmp_path):
    snapshot = model_registry_snapshot(bootstrap_manifest())
    repository = ModelLifecycleRepository(tmp_path / "lifecycle")
    publish_candidate(repository, snapshot, "existing-active")
    repository.replace_active(
        "existing-active",
        activated_at="2026-07-25T00:00:00+00:00",
        source="test",
        expected_revision=0,
    )
    service = TrainingLifecycleService(
        execution=_BoundedProductionExecution(snapshot),
        registry_provider=lambda: snapshot,
        repository=repository,
        publisher=build_candidate_publisher(repository),
    )
    data_path = write_mock_training_data(
        output_dir=tmp_path / "training-data",
        rows=18,
        seed=42,
    )

    service.start(TrainingRequest(
        run_id="production-integration",
        candidate_id="production-candidate",
        data_path=str(data_path),
    ))

    candidate = repository.read_candidate("production-candidate")
    result = json.loads(
        (candidate.path / "training_result.json").read_text(encoding="utf-8")
    )
    expected_identities = tuple(snapshot.target_presentation_order)
    result_identities = tuple(
        item["target_identity"] for item in result["targets"]
    )
    assert set(result_identities) == set(expected_identities)
    assert len(result_identities) == len(expected_identities)
    assert all(item["status"] == "complete" for item in result["targets"])
    assert all(
        set(item["metrics"]) >= {"r2", "mae", "rmse"}
        for item in result["targets"]
    )
    assert all(item["optuna"]["status"] == "used" for item in result["targets"])
    assert all(len(item["optuna"]["trials"]) == 1 for item in result["targets"])
    expected_rfecv = {
        target.identity: group.use_rfe
        for group in snapshot.groups
        for target in group.targets
    }
    assert all(
        (item["rfecv"]["status"] == "used")
        == expected_rfecv[item["target_identity"]]
        for item in result["targets"]
    )
    result_by_identity = {
        item["target_identity"]: item for item in result["targets"]
    }
    manifest_by_identity = {
        item.identity: item for item in candidate.manifest.targets
    }
    preprocessed = load_and_preprocess(data_path)
    for group in snapshot.groups:
        features, _targets = prepare_pipeline(
            preprocessed,
            {
                "name": group.name,
                "targets": [item.ml_name for item in group.targets],
                "use_rfe": group.use_rfe,
            },
            registry_snapshot=snapshot,
        )
        for target in group.targets:
            permitted = set(apply_target_policy(features.columns, target))
            trained = set(manifest_by_identity[target.identity].feature_names)
            selected = {
                item["feature"]
                for item in result_by_identity[target.identity]["rfecv"]["features"]
                if item["selected"]
            }
            assert trained <= permitted
            assert trained == selected
    quality = result["preprocessing"]["feature_data_quality"]
    assert all(
        item["target_usage_stage"] == "post_target_policy_pre_rfecv"
        for item in quality
    )
    assert {
        item["quality_collection_stage"] for item in quality
    } <= {
        "raw_training_input",
        "post_preprocessing_derived_features",
    }
    with (candidate.path / "analysis" / "target_metrics.csv").open(
        newline="", encoding="utf-8"
    ) as source:
        metrics_rows = list(csv.DictReader(source))
    workbook = load_workbook(
        candidate.path / "training_report.xlsx", data_only=True
    )
    assert [item["target_identity"] for item in metrics_rows] == list(
        result_identities
    )
    assert workbook["Target Metrics"].max_row == len(expected_identities) + 1
    assert repository.read_active().candidate_id == "existing-active"
