"""Mapping editor read-only snapshot export tests."""

import builtins
import json

from openpyxl import load_workbook

from core.mapping.editor_export import (
    READ_ONLY_SNAPSHOT_NOTICE,
    export_mapping_editor_snapshot_json,
    export_mapping_editor_snapshot_xlsx,
)
from core.mapping.editor_projection import project_runtime_mapping_to_editor_draft
from core.mapping.editor_validation import validate_mapping_editor_draft


MAPPING = {
    "idu": {"IDU-A": {"ID Volume": 1.25}},
    "evap_index": {},
    "odu": {"ODU-A": {"OD Volume": 2.5}},
    "compressor": {},
    "ref_type": {"R32": {}},
    "exp_type": {"EEV": {}},
    "odu_cascade": {
        "ODU-A": {
            "Available_Fins": ["F&T"],
            "Available_Pis": ["7"],
            "Available_Rows": ["1"],
        }
    },
    "cond_specs": {"ODU-A F&T 7 1": {"Cond Area": 3.5, "Cond Volume": 4.5}},
}


def test_json_snapshot_export_creates_read_only_review_file(tmp_path):
    draft = project_runtime_mapping_to_editor_draft(MAPPING)
    issues = validate_mapping_editor_draft(draft).issues
    export_file = tmp_path / "snapshot.json"

    result = export_mapping_editor_snapshot_json(draft, issues, export_file)

    payload = json.loads(export_file.read_text(encoding="utf-8"))
    assert result.success
    assert payload["snapshot_type"] == "mapping_editor_review_snapshot"
    assert payload["read_only"] is True
    assert payload["import_contract"] is False
    assert payload["groups"][0]["group"] == "IDU"
    assert payload["groups"][0]["rows"][0]["values"]["IDU"] == "IDU-A"
    assert payload["issues"] == []


def test_xlsx_snapshot_export_creates_read_only_workbook(tmp_path):
    draft = project_runtime_mapping_to_editor_draft(MAPPING)
    issues = validate_mapping_editor_draft(draft).issues
    export_file = tmp_path / "snapshot.xlsx"

    result = export_mapping_editor_snapshot_xlsx(draft, issues, export_file)

    workbook = load_workbook(export_file)
    assert result.success
    assert workbook.sheetnames == [
        "IDU",
        "Evap Index",
        "ODU",
        "Compressor",
        "Refrigerant",
        "Expansion",
        "ODU Cond Specs",
        "Issues",
        "Snapshot Info",
    ]
    idu_sheet = workbook["IDU"]
    assert [cell.value for cell in idu_sheet[1]] == [
        "IDU",
        "ID Volume",
        "Size",
        "source_key",
        "unresolved",
        "notes",
    ]
    assert idu_sheet["A2"].value == "IDU-A"
    assert idu_sheet["B2"].value == 1.25

    cond_sheet = workbook["ODU Cond Specs"]
    assert [cell.value for cell in cond_sheet[1]][:6] == [
        "ODU",
        "Fin Type",
        "Pi",
        "Row",
        "Cond Area",
        "Cond Volume",
    ]
    assert [cell.value for cell in cond_sheet[2]][:6] == ["ODU-A", "F&T", "7", "1", 3.5, 4.5]

    issues_sheet = workbook["Issues"]
    assert [cell.value for cell in issues_sheet[1]] == ["Level", "Group", "Row", "Field", "Message"]

    info_sheet = workbook["Snapshot Info"]
    assert info_sheet["B2"].value == READ_ONLY_SNAPSHOT_NOTICE
    assert info_sheet["B5"].value is False


def test_xlsx_snapshot_export_writes_issue_rows(tmp_path):
    mapping = {
        **MAPPING,
        "cond_specs": {"ODU-Z F&T 7 1": {"Cond Area": 3.5, "Cond Volume": 4.5}},
    }
    draft = project_runtime_mapping_to_editor_draft(mapping)
    issues = validate_mapping_editor_draft(draft).issues
    export_file = tmp_path / "snapshot-with-issues.xlsx"

    result = export_mapping_editor_snapshot_xlsx(draft, issues, export_file)

    workbook = load_workbook(export_file)
    assert result.success
    issue_values = [cell.value for cell in workbook["Issues"][2]]
    assert issue_values[0] == "error"
    assert issue_values[1] == "ODU Cond Specs"
    assert issue_values[4] == "Runtime cond_specs row is not matched by ODU cascade options."


def test_xlsx_snapshot_export_does_not_modify_runtime_mapping_json(tmp_path):
    mapping_file = tmp_path / "mapping.json"
    original_text = json.dumps(MAPPING, indent=2, ensure_ascii=False)
    mapping_file.write_text(original_text, encoding="utf-8")
    draft = project_runtime_mapping_to_editor_draft(MAPPING, source_label=str(mapping_file))

    result = export_mapping_editor_snapshot_xlsx(draft, (), tmp_path / "snapshot.xlsx")

    assert result.success
    assert mapping_file.read_text(encoding="utf-8") == original_text


def test_xlsx_snapshot_export_reports_missing_openpyxl(monkeypatch, tmp_path):
    draft = project_runtime_mapping_to_editor_draft(MAPPING)
    original_import = builtins.__import__

    def fail_openpyxl_import(name, *args, **kwargs):
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError("blocked for test")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fail_openpyxl_import)

    result = export_mapping_editor_snapshot_xlsx(draft, (), tmp_path / "snapshot.xlsx")

    assert not result.success
    assert "openpyxl is required" in result.message
