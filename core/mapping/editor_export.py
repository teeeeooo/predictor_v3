"""Export read-only mapping editor review snapshots."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from core.mapping.editor_model import MappingEditorDraft
from core.mapping.entity_model import MappingValidationError

READ_ONLY_SNAPSHOT_NOTICE = "Read-only review snapshot. This file is not an import contract."
_ISSUE_HEADERS = ("Level", "Group", "Row", "Field", "Message")
_SUPPORT_HEADERS = ("source_key", "unresolved", "notes")


@dataclass(frozen=True)
class MappingEditorExportResult:
    """Result of exporting a mapping editor review snapshot."""

    success: bool
    path: Path
    message: str = ""


def mapping_editor_snapshot_payload(
    draft: MappingEditorDraft,
    issues: tuple[MappingValidationError, ...] = (),
) -> dict[str, Any]:
    """Return a read-only user-facing snapshot payload."""
    return {
        "snapshot_type": "mapping_editor_review_snapshot",
        "read_only": True,
        "import_contract": False,
        "source": draft.source_label,
        "groups": [
            {
                "group": group.label,
                "key": group.group_key,
                "columns": list(group.columns),
                "rows": [
                    {
                        "values": {column: row.value_for(column, "") for column in group.columns},
                        "source_key": row.source_key,
                        "unresolved": row.unresolved,
                        "notes": row.notes,
                    }
                    for row in group.rows
                ],
            }
            for group in draft.groups
        ],
        "issues": [
            {
                "level": issue.severity,
                "group": issue.entity_key,
                "row": issue.row_key,
                "field": issue.attribute_key or issue.field,
                "message": issue.message,
            }
            for issue in issues
        ],
    }


def export_mapping_editor_snapshot_json(
    draft: MappingEditorDraft,
    issues: tuple[MappingValidationError, ...],
    destination: str | Path,
) -> MappingEditorExportResult:
    """Write a pretty JSON read-only review snapshot."""
    path = Path(destination)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        payload = mapping_editor_snapshot_payload(draft, issues)
        with path.open("w", encoding="utf-8") as json_file:
            json.dump(payload, json_file, indent=2, ensure_ascii=False)
            json_file.write("\n")
    except Exception as exc:
        return MappingEditorExportResult(
            success=False,
            path=path,
            message=f"Export failed: {exc}",
        )
    return MappingEditorExportResult(
        success=True,
        path=path,
        message=f"Exported read-only review snapshot to {path}.",
    )


def export_mapping_editor_snapshot_xlsx(
    draft: MappingEditorDraft,
    issues: tuple[MappingValidationError, ...],
    destination: str | Path,
) -> MappingEditorExportResult:
    """Write a read-only XLSX review snapshot."""
    path = Path(destination)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return MappingEditorExportResult(
            success=False,
            path=path,
            message="Export failed: openpyxl is required for XLSX export.",
        )

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        payload = mapping_editor_snapshot_payload(draft, issues)
        workbook = Workbook()
        workbook.remove(workbook.active)

        for group in payload["groups"]:
            worksheet = workbook.create_sheet(_safe_sheet_title(workbook, group["group"]))
            headers = [*group["columns"], *_SUPPORT_HEADERS]
            rows = [
                [
                    *(_xlsx_cell_value(row["values"].get(column, "")) for column in group["columns"]),
                    _xlsx_cell_value(row["source_key"]),
                    _xlsx_cell_value(row["unresolved"]),
                    _xlsx_cell_value(row["notes"]),
                ]
                for row in group["rows"]
            ]
            _write_table(worksheet, headers, rows, Font)

        issues_sheet = workbook.create_sheet("Issues")
        issue_rows = [
            [issue["level"], issue["group"], issue["row"], issue["field"], issue["message"]]
            for issue in payload["issues"]
        ]
        _write_table(issues_sheet, _ISSUE_HEADERS, issue_rows, Font)

        info_sheet = workbook.create_sheet("Snapshot Info")
        _write_table(
            info_sheet,
            ("Field", "Value"),
            [
                ("Notice", READ_ONLY_SNAPSHOT_NOTICE),
                ("Snapshot Type", payload["snapshot_type"]),
                ("Read Only", payload["read_only"]),
                ("Import Contract", payload["import_contract"]),
                ("Source", payload["source"]),
            ],
            Font,
        )

        workbook.save(path)
    except Exception as exc:
        return MappingEditorExportResult(
            success=False,
            path=path,
            message=f"Export failed: {exc}",
        )
    return MappingEditorExportResult(
        success=True,
        path=path,
        message=f"Exported read-only review snapshot to {path}.",
    )


def _safe_sheet_title(workbook: Any, title: str) -> str:
    candidate = "".join("_" if character in "[]:*?/\\" else character for character in title).strip()
    candidate = candidate[:31] or "Sheet"
    existing = set(workbook.sheetnames)
    if candidate not in existing:
        return candidate

    base = candidate[:28]
    index = 2
    while f"{base}_{index}" in existing:
        index += 1
    return f"{base}_{index}"


def _write_table(worksheet: Any, headers: tuple[str, ...] | list[str], rows: list[list[Any]], font_type: Any) -> None:
    worksheet.append(list(headers))
    for cell in worksheet[1]:
        cell.font = font_type(bold=True)
    for row in rows:
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"
    if headers:
        worksheet.auto_filter.ref = worksheet.dimensions
    _fit_columns(worksheet)


def _fit_columns(worksheet: Any) -> None:
    for column_cells in worksheet.columns:
        header = column_cells[0]
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        worksheet.column_dimensions[header.column_letter].width = min(max(max_length + 2, 10), 48)


def _xlsx_cell_value(value: Any) -> str | int | float | bool:
    if value is None:
        return ""
    if isinstance(value, str | int | float | bool):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)
