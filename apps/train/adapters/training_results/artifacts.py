"""Filesystem adapter deriving JSON, CSV, and XLSX from one result contract."""

from __future__ import annotations

import csv
import hashlib
import json
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from apps.common.model_lifecycle.candidate_contracts import CandidateArtifactReference
from apps.train.application.training_results import TrainingAnalysisResult
from .tables import result_tables, workbook_tables


class TrainingResultArtifactWriter:
    """Write and verify the required Candidate-owned analysis artifacts."""

    def write(
        self, staging: Path, result: TrainingAnalysisResult
    ) -> tuple[CandidateArtifactReference, ...]:
        analysis = staging / "analysis"
        analysis.mkdir(exist_ok=True)
        payload = result.to_payload()
        _write_json(staging / "training_result.json", payload)
        tables = result_tables(result)
        for filename, rows in tables.items():
            _write_csv(analysis / filename, rows)
        _write_workbook(staging / "training_report.xlsx", result, tables)
        self.verify(staging, result)
        return tuple(
            CandidateArtifactReference(
                path=str(item["path"]),
                sha256=artifact_sha256(staging / str(item["path"])),
                category=str(item["category"]),
                required=item["required"],
            )
            for item in result.artifacts
        )

    def verify(self, staging: Path, result: TrainingAnalysisResult) -> None:
        payload = json.loads(
            (staging / "training_result.json").read_text(encoding="utf-8")
        )
        loaded = TrainingAnalysisResult.from_payload(payload)
        normalized_source = json.loads(
            json.dumps(result.to_payload(), ensure_ascii=False, sort_keys=True)
        )
        normalized_loaded = json.loads(
            json.dumps(loaded.to_payload(), ensure_ascii=False, sort_keys=True)
        )
        if normalized_loaded != normalized_source:
            raise ValueError("training result JSON does not match source contract")
        tables = result_tables(result)
        for filename, expected in tables.items():
            with (staging / "analysis" / filename).open(
                newline="", encoding="utf-8"
            ) as source:
                observed = list(csv.DictReader(source))
            if observed != _string_rows(expected):
                raise ValueError(f"training result CSV mismatch: {filename}")
        workbook = load_workbook(staging / "training_report.xlsx", data_only=True)
        expected_sheets = tuple(workbook_tables(result, tables))
        if tuple(workbook.sheetnames) != expected_sheets:
            raise ValueError("training report XLSX sheet contract mismatch")
        for title, rows in workbook_tables(result, tables).items():
            sheet = workbook[title]
            headers = list(rows[0]) if rows else ["status"]
            observed_headers = [
                sheet.cell(1, column).value
                for column in range(1, len(headers) + 1)
            ]
            if observed_headers != headers:
                raise ValueError(f"training report XLSX header mismatch: {title}")
            observed_rows = [
                {
                    header: sheet.cell(row_index, column_index).value
                    for column_index, header in enumerate(headers, start=1)
                }
                for row_index in range(2, sheet.max_row + 1)
            ]
            expected_rows = [
                {key: _cell_value(value) for key, value in row.items()}
                for row in rows
            ]
            if not _workbook_rows_equal(observed_rows, expected_rows):
                raise ValueError(f"training report XLSX value mismatch: {title}")


class MinimalTrainingResultEvidenceWriter:
    """Persist terminal JSON without depending on CSV/XLSX generation."""

    @staticmethod
    def write(staging: Path, result: TrainingAnalysisResult) -> None:
        _write_json(staging / "training_result.json", result.to_payload())


def artifact_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_workbook(
    path: Path,
    result: TrainingAnalysisResult,
    tables: dict[str, list[dict[str, Any]]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in workbook_tables(result, tables).items():
        sheet = workbook.create_sheet(title)
        _append_rows(sheet, rows)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)


def _append_rows(sheet, rows: list[dict[str, Any]]) -> None:  # noqa: ANN001
    headers = list(rows[0]) if rows else ("status",)
    sheet.append(headers)
    for row in rows:
        sheet.append([_cell_value(row.get(header, "")) for header in headers])


def _cell_value(value: Any) -> Any:
    if value == "":
        return None
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _workbook_rows_equal(
    observed: list[dict[str, Any]],
    expected: list[dict[str, Any]],
) -> bool:
    if len(observed) != len(expected):
        return False
    for observed_row, expected_row in zip(observed, expected):
        if observed_row.keys() != expected_row.keys():
            return False
        for key, expected_value in expected_row.items():
            observed_value = observed_row[key]
            if isinstance(expected_value, float) and isinstance(
                observed_value, (int, float)
            ):
                if abs(float(observed_value) - expected_value) > 1e-12:
                    return False
            elif observed_value != expected_value:
                return False
    return True


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    headers = list(rows[0]) if rows else ("status",)
    with path.open("w", newline="", encoding="utf-8") as target:
        writer = csv.DictWriter(target, fieldnames=headers)
        writer.writeheader()
        writer.writerows(_string_rows(rows))


def _string_rows(rows: Iterable[dict[str, Any]]) -> list[dict[str, str]]:
    return [
        {
            key: (
                json.dumps(value, ensure_ascii=False, sort_keys=True)
                if isinstance(value, (dict, list, tuple))
                else str(value)
            )
            for key, value in row.items()
        }
        for row in rows
    ]
